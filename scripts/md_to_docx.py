"""Convert the tutorial markdown (headings, bullets, numbered lists, bold, inline code, images) to .docx.

    python scripts/md_to_docx.py tutorial/growth_chart_tutorial_zh.md tutorial/growth_chart_tutorial_zh.docx
"""
import os
import re
import sys
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.oxml.ns import qn

src, out = sys.argv[1], sys.argv[2]
base = os.path.dirname(os.path.abspath(src))
doc = Document()

# fonts: Latin + East Asian
style = doc.styles["Normal"]
style.font.name = "Calibri"
style.font.size = Pt(11)
style.element.rPr.rFonts.set(qn("w:eastAsia"), "PingFang SC")
for s in ("Heading 1", "Heading 2", "Heading 3", "Title"):
    doc.styles[s].element.rPr.rFonts.set(qn("w:eastAsia"), "PingFang SC")
    doc.styles[s].font.color.rgb = RGBColor(0x1F, 0x5F, 0xA8)
for sec in doc.sections:
    sec.left_margin = sec.right_margin = Inches(1)

INLINE = re.compile(r"(\*\*.+?\*\*|`.+?`)")


def add_runs(par, text):
    for piece in INLINE.split(text):
        if not piece:
            continue
        if piece.startswith("**"):
            r = par.add_run(piece[2:-2]); r.bold = True
        elif piece.startswith("`"):
            r = par.add_run(piece[1:-1]); r.font.name = "Menlo"; r.font.size = Pt(10)
        else:
            par.add_run(piece)


lines = open(src, encoding="utf-8").read().splitlines()
i = 0
while i < len(lines):
    ln = lines[i]
    if ln.startswith("# "):
        doc.add_heading(ln[2:], level=0)
    elif ln.startswith("## "):
        doc.add_heading(ln[3:], level=1)
    elif ln.startswith("### "):
        doc.add_heading(ln[4:], level=2)
    elif m := re.match(r"!\[.*?\]\((.+?)\)", ln.strip()):
        doc.add_picture(os.path.join(base, m.group(1)), width=Inches(6.3))
    elif re.match(r"^\s*[-*] ", ln):
        indent = len(ln) - len(ln.lstrip())
        p = doc.add_paragraph(style="List Bullet 2" if indent >= 2 else "List Bullet")
        add_runs(p, re.sub(r"^\s*[-*] ", "", ln))
    elif re.match(r"^\s*\d+\. ", ln):
        p = doc.add_paragraph(style="List Number")
        add_runs(p, re.sub(r"^\s*\d+\. ", "", ln))
    elif ln.startswith("|"):
        rows = []
        while i < len(lines) and lines[i].startswith("|"):
            cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
            if not all(re.fullmatch(r":?-+:?", c) for c in cells):
                rows.append(cells)
            i += 1
        i -= 1
        t = doc.add_table(rows=len(rows), cols=len(rows[0])); t.style = "Light Grid Accent 1"
        for r_i, row in enumerate(rows):
            for c_i, cell in enumerate(row[:len(rows[0])]):
                par = t.cell(r_i, c_i).paragraphs[0]
                add_runs(par, cell.replace("\\|", "|"))
                for run in par.runs:
                    run.font.size = Pt(9.5)
                    if r_i == 0:
                        run.bold = True
    elif ln.startswith("```"):
        j = i + 1
        while j < len(lines) and not lines[j].startswith("```"):
            j += 1
        p = doc.add_paragraph()
        r = p.add_run("\n".join(lines[i + 1:j])); r.font.name = "Menlo"; r.font.size = Pt(9.5)
        i = j
    elif ln.strip() == "":
        pass
    else:
        # merge wrapped paragraph lines
        buf = [ln]
        while i + 1 < len(lines) and lines[i + 1].strip() and not re.match(r"^(#|\s*[-*] |\s*\d+\. |!\[|```)", lines[i + 1]):
            i += 1; buf.append(lines[i])
        p = doc.add_paragraph()
        add_runs(p, " ".join(s.strip() for s in buf))
    i += 1

doc.save(out)
print(out, os.path.getsize(out))
